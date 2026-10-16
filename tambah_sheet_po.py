"""Tambah sheet "PO Diskominfo" (ringkasan per NO. PO) + "Rincian PO" (per item) dari teks data PO.

Format teks: blok dipisah baris kosong. Tiap blok berisi:
    4600071327                          <- NO. PO
    PA/ACT/2002/0254/TER                <- NO. PA/ACT
    DINAS ... KOTA SINGKAWANG - Lokasi  <- pelanggan (+ nama lokasi)
    Jl. ...                             <- alamat
    102020B06767                        <- ID layanan
    11.498.673                          <- estimasi harga
    3003571<TAB>Survey ...<TAB>1<TAB>LOT<TAB>821.630   <- rincian (kode boleh kosong)

Pemakaian:
    python tambah_sheet_po.py <workbook.xlsx> <data_po.txt> [output.xlsx]
"""
import re
import sys
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

from rapikan_sheet2 import BORDER, FMT_RUPIAH, angka

RINGKASAN = "PO Diskominfo"
RINCIAN = "Rincian PO"

RE_PO = re.compile(r"^46\d{8}$")
RE_PA = re.compile(r"^PA/ACT/")
RE_ID = re.compile(r"^\d{4,6}B[A-Z]?\d{5,6}$")
RE_TOTAL = re.compile(r"^\d{1,3}(?:\.\d{3})+$")
RE_ITEM = re.compile(
    r"^(?:(\d{7})\s+)?(.+?)\s+(\d+(?:[.,]\d+)*)\s+([A-Za-z]+)\s+(\d{1,3}(?:\.\d{3})*)$")
RE_PELANGGAN = re.compile(r"^(DINAS\b.*?\bKOTA SINGKAWANG)\b[\s|\-]*(.*)$")


def pecah_lokasi(teks):
    """'Kantor X - Jl. ...' atau 'Kantor X Jl. ...' -> ('Kantor X', 'Jl. ...')."""
    if " - " in teks:
        lok, alamat = teks.split(" - ", 1)
    else:
        m = re.search(r"\s(?=Jl\.?\s)", teks)
        lok, alamat = (teks[:m.start()], teks[m.end():]) if m else (teks, "")
    if re.match(r"Jl\.?\s", lok):  # baris ini hanya alamat
        return "", teks.strip()
    return lok.strip(), alamat.strip()


def pecah_blok(blok):
    data = {"po": None, "pa": None, "id": None, "estimasi": None, "items": []}
    teks = []
    for baris in blok.splitlines():
        baris = baris.strip()
        if not baris:
            continue
        if RE_PO.match(baris):
            data["po"] = int(baris)
        elif RE_PA.match(baris):
            data["pa"] = baris
        elif RE_ID.match(baris):
            data["id"] = baris
        elif RE_TOTAL.match(baris):
            data["estimasi"] = angka(baris)
        elif m := RE_ITEM.match(baris):
            kode, uraian, vol, satuan, harga = m.groups()
            data["items"].append([int(kode) if kode else None, uraian.strip(),
                                  angka(vol), satuan, angka(harga)])
        else:
            teks.append(baris)

    m = RE_PELANGGAN.match(teks[0]) if teks else None
    pelanggan, sisa = m.groups() if m else ("", teks[0] if teks else "")
    lain = " ".join(teks[1:])
    if sisa:
        lok, alamat = pecah_lokasi(sisa)
        alamat = alamat or lain
        # 'CCTV Sekip Lama - Jl. Pemuda' -> nama lokasi tidak perlu diulang di alamat
        if lok and alamat.startswith(lok + " - "):
            alamat = alamat[len(lok) + 3:]
    else:
        lok, alamat = pecah_lokasi(lain)
    data["pelanggan"] = pelanggan
    data["lokasi"] = (lok or pelanggan).upper()
    data["alamat"] = alamat
    return data


def tulis_header(ws, headers, widths, font):
    for c, (judul, lebar) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, c, judul)
        cell.font = copy(font)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = lebar


def tulis_total(ws, row, label_col, cols, font):
    ws.cell(row, label_col, "TOTAL").font = copy(font)
    ws.cell(row, label_col).alignment = Alignment(horizontal="right")
    for col in cols:
        cell = ws.cell(row, ord(col) - 64, f"=SUM({col}2:{col}{row - 1})")
        cell.font = copy(font)
        cell.number_format = FMT_RUPIAH
        cell.border = Border(top=Side(style="thin"), bottom=Side(style="double"))


def main(src, data_path, dst):
    with open(data_path, encoding="utf-8") as f:
        bloks = [pecah_blok(b) for b in re.split(r"\n\s*\n", f.read()) if b.strip()]

    wb = load_workbook(src)
    for nama in (RINGKASAN, RINCIAN):
        if nama in wb.sheetnames:
            del wb[nama]
    ref = wb["Sheet2"]
    font_head = copy(ref.cell(1, 8).font)
    font_body = copy(ref.cell(2, 5).font)
    font_bold = copy(font_body)
    font_bold.b = True

    # Kode yang kosong diisi dari uraian yang sama (data ini + sheet Rincian RAB).
    kode_by_uraian = {}
    if "Rincian RAB" in wb.sheetnames:
        for kode, uraian in wb["Rincian RAB"].iter_rows(min_row=2, min_col=5, max_col=6, values_only=True):
            if kode and uraian:
                kode_by_uraian[uraian] = kode
    for b in bloks:
        for item in b["items"]:
            if item[0]:
                kode_by_uraian[item[1]] = item[0]

    # 1. Rincian PO: satu baris per item.
    wr = wb.create_sheet(RINCIAN)
    tulis_header(wr, ["NO", "NO. PO", "NAMA LOKASI", "KODE", "URAIAN PEKERJAAN",
                      "VOLUME", "SATUAN", "HARGA SATUAN", "JUMLAH"],
                 [6, 13, 40, 10, 42, 10, 9, 15, 16], font_head)
    formats = {2: "0", 4: "0", 6: "#,##0", 8: FMT_RUPIAH, 9: FMT_RUPIAH}
    i = 2
    for b in bloks:
        for kode, uraian, vol, satuan, harga in b["items"]:
            kode = kode or kode_by_uraian.get(uraian)
            values = [i - 1, b["po"], b["lokasi"], kode, uraian, vol, satuan, harga, f"=F{i}*H{i}"]
            for c, v in enumerate(values, start=1):
                cell = wr.cell(i, c, v)
                cell.font = copy(font_body)
                cell.border = BORDER
                if c in formats:
                    cell.number_format = formats[c]
                if c in (1, 7):
                    cell.alignment = Alignment(horizontal="center")
            i += 1
    last_item = i - 1
    tulis_total(wr, i, 8, "I", font_bold)
    wr.freeze_panes = "A2"
    wr.auto_filter.ref = f"A1:I{last_item}"

    # 2. PO Diskominfo: satu baris per NO. PO, dicocokkan ke rincian lewat NO. PO.
    ws = wb.create_sheet(RINGKASAN)
    tulis_header(ws, ["NO", "NO. PO", "NO. PA/ACT", "NAMA LOKASI", "PELANGGAN", "ALAMAT",
                      "ID LAYANAN", "ESTIMASI HARGA", "TOTAL RINCIAN", "SELISIH"],
                 [5, 13, 22, 40, 34, 60, 15, 16, 16, 14], font_head)
    jumlah = f"'{RINCIAN}'!$I$2:$I${last_item}"
    po_col = f"'{RINCIAN}'!$B$2:$B${last_item}"
    for r, b in enumerate(bloks, start=2):
        values = [r - 1, b["po"], b["pa"], b["lokasi"], b["pelanggan"], b["alamat"], b["id"],
                  b["estimasi"],
                  f'=IF($B{r}="",0,SUMIFS({jumlah},{po_col},$B{r}))',
                  f'=IF(I{r}=0,"",H{r}-I{r})']
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.font = copy(font_body)
            cell.border = BORDER
        ws.cell(r, 1).alignment = Alignment(horizontal="center")
        ws.cell(r, 2).number_format = "0"
        ws.cell(r, 8).number_format = FMT_RUPIAH
        ws.cell(r, 9).number_format = '#,##0;-#,##0;"-"'
        ws.cell(r, 10).number_format = "#,##0;[Red]-#,##0;0"
    tulis_total(ws, len(bloks) + 2, 7, "HIJ", font_bold)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(bloks) + 1}"

    # Ringkasan diletakkan sebelum rinciannya.
    wb.move_sheet(ws, offset=-1)

    try:
        wb.save(dst)
    except PermissionError:
        sys.exit(f"Gagal menyimpan: {dst} sedang dibuka di Excel. Tutup dulu lalu jalankan lagi.")
    print(f"{len(bloks)} PO, {last_item - 1} baris rincian -> {dst}")


if __name__ == "__main__":
    if len(sys.argv) not in (3, 4):
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) == 4 else sys.argv[1])
