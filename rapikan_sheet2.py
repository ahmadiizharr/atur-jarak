"""Rapikan Sheet2: pecah rincian RAB (teks berkode & bertitik) jadi tabel angka yang bisa di-SUM.

Kolom J/K/L di Sheet2 berisi rincian pekerjaan dalam bentuk teks, misalnya:
    3003571<TAB>Survey jalur kabel incl.Spotting GPS Blk<TAB>1<TAB>LOT<TAB>821.630
Script ini memindahkan setiap baris rincian ke sheet "Rincian RAB" (kode, uraian,
volume, satuan, harga satuan sebagai angka + rumus JUMLAH), lalu mengganti kolom
J/K/L di Sheet2 dengan TOTAL RINCIAN, SELISIH, dan KETERANGAN, plus baris TOTAL.

Pemakaian:
    python rapikan_sheet2.py <input.xlsx> <output.xlsx>
"""
import re
import sys
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

SHEET = "Sheet2"
RINCIAN = "Rincian RAB"
LIST_COLS = (10, 11, 12)  # J, K, L
COL_ESTIMASI = 8  # H
COL_PO = 9  # I
FMT_RUPIAH = "#,##0"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def angka(teks):
    """'1.089.630' -> 1089630, '1,5' -> 1.5 (format angka Indonesia)."""
    v = float(str(teks).strip().replace(".", "").replace(",", "."))
    return int(v) if v.is_integer() else v


def pecah_baris(baris):
    """Satu baris rincian -> (kode, uraian, volume, satuan, harga) atau None."""
    bagian = [b.strip() for b in re.split(r"\t| {2,}", baris.strip()) if b.strip()]
    kode = None
    if bagian and re.fullmatch(r"\d{7}", bagian[0]):
        kode = int(bagian.pop(0))
    if len(bagian) != 4:
        return None
    uraian, vol, satuan, harga = bagian
    try:
        return kode, uraian, angka(vol), satuan, angka(harga)
    except ValueError:
        return None


def pecah_rincian(nilai):
    """Teks rincian -> list item, atau None kalau sel ini bukan rincian."""
    if not isinstance(nilai, str) or not re.search(r"\t| {2,}", nilai):
        return None
    items = []
    for baris in nilai.split("\n"):
        if not baris.strip():
            continue
        item = pecah_baris(baris)
        if item is None:
            return None
        items.append(item)
    return items or None


def mirip_po(nilai):
    if isinstance(nilai, (int, float)):
        return True
    return isinstance(nilai, str) and nilai.strip().isdigit()


def no_po(ws, row, col):
    """No. PO untuk rincian di kolom `col`: sel di kirinya kalau berisi nomor,
    selain itu bagian ke-n dari kolom I yang dipisah '/' (J=1, K=2, L=3)."""
    kiri = ws.cell(row, col - 1).value
    if mirip_po(kiri):
        return int(kiri)
    bagian = [b.strip() for b in str(ws.cell(row, COL_PO).value or "").split("/")]
    idx = col - LIST_COLS[0]
    if idx < len(bagian) and bagian[idx].isdigit():
        return int(bagian[idx])
    return None


def lokasi(alamat):
    """'[311202004340] PUSKESMAS ... - JL ...' -> 'PUSKESMAS ... - JL ...'"""
    return re.sub(r"^\[\d+\]\s*", "", str(alamat or "")).strip()


def main(src, dst):
    wb = load_workbook(src)
    ws = wb[SHEET]

    last = max(r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value is not None)

    # 1. Kumpulkan semua rincian + keterangan (nomor PO yang ada di J/K/L).
    rows = []  # [service_id, lokasi, po, kode, uraian, vol, satuan, harga]
    keterangan = {}
    for r in range(2, last + 1):
        catatan = []
        for c in LIST_COLS:
            nilai = ws.cell(r, c).value
            items = pecah_rincian(nilai)
            if items:
                po = no_po(ws, r, c)
                for kode, uraian, vol, satuan, harga in items:
                    rows.append([ws.cell(r, 2).value, lokasi(ws.cell(r, 5).value), po,
                                 kode, uraian, vol, satuan, harga])
            elif mirip_po(nilai):
                catatan.append(f"No. PO: {int(nilai)}")
        if catatan:
            keterangan[r] = "; ".join(catatan)

    # Kode yang kosong diisi dari uraian yang sama di baris lain.
    kode_by_uraian = {row[4]: row[3] for row in rows if row[3] is not None}
    for row in rows:
        if row[3] is None:
            row[3] = kode_by_uraian.get(row[4])

    # 2. Sheet "Rincian RAB".
    if RINCIAN in wb.sheetnames:
        del wb[RINCIAN]
    wr = wb.create_sheet(RINCIAN)
    font_head = copy(ws.cell(1, COL_ESTIMASI).font)
    font_body = copy(ws.cell(2, 5).font)
    font_bold = copy(font_body)
    font_bold.b = True

    headers = ["NO", "SERVICE ID", "LOKASI", "NO. PO", "KODE", "URAIAN PEKERJAAN",
               "VOLUME", "SATUAN", "HARGA SATUAN", "JUMLAH"]
    widths = [6, 15, 60, 13, 10, 42, 10, 9, 15, 16]
    for c, (judul, lebar) in enumerate(zip(headers, widths), start=1):
        cell = wr.cell(1, c, judul)
        cell.font = copy(font_head)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        wr.column_dimensions[cell.column_letter].width = lebar

    formats = {2: "0", 4: "0", 5: "0", 7: "#,##0", 9: FMT_RUPIAH, 10: FMT_RUPIAH}
    for i, row in enumerate(rows, start=2):
        values = [i - 1, *row, f"=G{i}*I{i}"]
        for c, v in enumerate(values, start=1):
            cell = wr.cell(i, c, v)
            cell.font = copy(font_body)
            cell.border = BORDER
            if c in formats:
                cell.number_format = formats[c]
            if c in (1, 8):
                cell.alignment = Alignment(horizontal="center")
    last_item = len(rows) + 1
    total_row = last_item + 1
    wr.cell(total_row, 9, "TOTAL").font = copy(font_bold)
    wr.cell(total_row, 9).alignment = Alignment(horizontal="right")
    cell = wr.cell(total_row, 10, f"=SUM(J2:J{last_item})")
    cell.font = copy(font_bold)
    cell.number_format = FMT_RUPIAH
    cell.border = BORDER
    wr.freeze_panes = "A2"
    wr.auto_filter.ref = f"A1:J{last_item}"

    # 3. Sheet2: J/K/L jadi TOTAL RINCIAN, SELISIH, KETERANGAN + baris TOTAL.
    for r in range(1, last + 1):
        for c in LIST_COLS:
            ws.cell(r, c).value = None
    for c, (judul, lebar) in zip(LIST_COLS, [("TOTAL RINCIAN", 17), ("SELISIH", 15), ("KETERANGAN", 24)]):
        cell = ws.cell(1, c, judul)
        cell.font = copy(font_head)
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = lebar

    jumlah = f"'{RINCIAN}'!$J$2:$J${last_item}"
    service = f"'{RINCIAN}'!$B$2:$B${last_item}"
    for r in range(2, last + 1):
        j = ws.cell(r, 10, f"=SUMIFS({jumlah},{service},$B{r})")
        j.number_format = '#,##0;-#,##0;"-"'
        k = ws.cell(r, 11, f'=IF(J{r}=0,"",H{r}-J{r})')
        k.number_format = "#,##0;[Red]-#,##0;0"
        for cell in (j, k):
            cell.font = copy(font_body)
            cell.alignment = Alignment(horizontal="right")
        if r in keterangan:
            ws.cell(r, 12, keterangan[r]).font = copy(font_body)

    t = last + 1
    ws.cell(t, 7, "TOTAL").font = copy(font_bold)
    ws.cell(t, 7).alignment = Alignment(horizontal="right")
    for c, col in ((8, "H"), (10, "J"), (11, "K")):
        cell = ws.cell(t, c, f"=SUM({col}2:{col}{last})")
        cell.font = copy(font_bold)
        cell.number_format = FMT_RUPIAH
        cell.alignment = Alignment(horizontal="right")
        cell.border = Border(top=Side(style="thin"), bottom=Side(style="double"))

    wb.save(dst)
    print(f"{len(rows)} baris rincian dari {len({r[0] for r in rows})} lokasi -> {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
